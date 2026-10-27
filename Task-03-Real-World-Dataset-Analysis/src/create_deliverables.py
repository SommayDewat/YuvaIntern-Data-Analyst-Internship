from pathlib import Path
import pandas as pd
import nbformat as nbf
from docx import Document
from docx.shared import Inches, Pt

ROOT = Path(__file__).resolve().parents[1]
RESULTS = ROOT / "results"
REPORTS = ROOT / "reports"
NOTEBOOKS = ROOT / "notebooks"

df = pd.read_csv(RESULTS / "cleaned_seoul_bike_data.csv", parse_dates=["Date"])
rentals = df["Rented_Bike_Count"]
hourly = df.groupby("Hour")["Rented_Bike_Count"].mean()
seasonal = df.groupby("Seasons")["Rented_Bike_Count"].mean().sort_values(ascending=False)
peak = df.loc[rentals.idxmax()]

# Create Jupyter Notebook
nb = nbf.v4.new_notebook()
nb["metadata"] = {
    "kernelspec": {"display_name": "Python 3", "language": "python", "name": "python3"}
}
nb["cells"] = [
    nbf.v4.new_markdown_cell(
        "# Seoul Bike Sharing Demand Analysis\n"
        "YuvaIntern Data Analyst Internship — Task 3\n\n"
        "Dataset source: UCI Machine Learning Repository\n"
        "https://archive.ics.uci.edu/dataset/560/seoul%2Bbike%2Bsharing%2Bdemand"
    ),
    nbf.v4.new_code_cell(
        "from pathlib import Path\nimport pandas as pd\nimport numpy as np\n"
        "import matplotlib.pyplot as plt\n"
        "df = pd.read_csv(Path('results/cleaned_seoul_bike_data.csv'), parse_dates=['Date'])\n"
        "df['Day_Type'] = np.where(df['Date'].dt.dayofweek >= 5, 'Weekend', 'Weekday')\n"
        "print('Shape:', df.shape)\ndisplay(df.head())\ndisplay(df.describe().T)"
    ),
    nbf.v4.new_markdown_cell("## Data quality"),
    nbf.v4.new_code_cell(
        "print('Missing values:', int(df.isna().sum().sum()))\n"
        "print('Duplicate rows:', int(df.duplicated().sum()))\n"
        "print('Duplicate date-hour keys:', int(df.duplicated(['Date','Hour']).sum()))"
    ),
    nbf.v4.new_markdown_cell("## Hourly demand"),
    nbf.v4.new_code_cell(
        "hourly = df.groupby('Hour')['Rented_Bike_Count'].mean()\n"
        "display(hourly.round(2).to_frame('average_rentals'))\n"
        "hourly.plot(figsize=(9,4), marker='o', title='Average Rentals by Hour')\n"
        "plt.xlabel('Hour'); plt.ylabel('Average rentals'); plt.grid(alpha=.3)\n"
        "plt.tight_layout(); plt.show()"
    ),
    nbf.v4.new_markdown_cell("## Seasonal demand"),
    nbf.v4.new_code_cell(
        "seasonal = df.groupby('Seasons')['Rented_Bike_Count'].mean().sort_values(ascending=False)\n"
        "display(seasonal.round(2).to_frame('average_rentals'))\n"
        "seasonal.plot(kind='bar', title='Average Rentals by Season')\n"
        "plt.ylabel('Average rentals/hour'); plt.xticks(rotation=0)\n"
        "plt.tight_layout(); plt.show()"
    ),
    nbf.v4.new_markdown_cell("## Weekday versus weekend"),
    nbf.v4.new_code_cell(
        "display(df.groupby('Day_Type')['Rented_Bike_Count'].agg(['mean','median','sum']).round(2))"
    ),
    nbf.v4.new_markdown_cell("## Temperature relationship"),
    nbf.v4.new_code_cell(
        "print(df.select_dtypes(include='number').corr(numeric_only=True)['Rented_Bike_Count'].sort_values(ascending=False))\n"
        "plt.scatter(df['TemperatureC'], df['Rented_Bike_Count'], alpha=.25)\n"
        "plt.xlabel('Temperature (°C)'); plt.ylabel('Bike rentals')\n"
        "plt.title('Temperature versus Bike Rentals'); plt.show()"
    ),
    nbf.v4.new_markdown_cell(
        "## Conclusions\n"
        "- The dataset contains 8,760 hourly records.\n"
        "- Total rentals are 6,172,314.\n"
        "- Average hourly demand peaks at 18:00.\n"
        "- Summer has the highest seasonal average; winter has the lowest.\n"
        "- Correlation shows association, not causation.\n"
        "- Potential outliers should be investigated rather than automatically removed."
    )
]

nbf.write(nb, NOTEBOOKS / "01_seoul_bike_demand_analysis.ipynb")

# Create Word report
doc = Document()
doc.styles["Normal"].font.name = "Arial"
doc.styles["Normal"].font.size = Pt(10)
doc.add_heading("Real-World Dataset Analysis Project", 0)
doc.add_paragraph("Seoul Bike Sharing Demand — YuvaIntern Task 3")

doc.add_heading("1. Dataset and Objective", level=1)
doc.add_paragraph(
    "This project explores the Seoul Bike Sharing Demand dataset from the UCI Machine "
    "Learning Repository. The objective is to inspect data quality, analyse hourly and "
    "seasonal demand, examine weather relationships, and communicate findings."
)
doc.add_paragraph(
    "Source: https://archive.ics.uci.edu/dataset/560/seoul%2Bbike%2Bsharing%2Bdemand"
)
doc.add_paragraph(
    "The dataset includes rental counts, date, hour, weather measurements, season, "
    "holiday status, and system functioning status. The source dataset is distributed "
    "under CC BY 4.0; attribution should be retained."
)

doc.add_heading("2. Data Quality", level=1)
doc.add_paragraph(f"Records: {len(df):,}")
doc.add_paragraph(f"Missing values: {int(df.isna().sum().sum())}")
doc.add_paragraph(f"Exact duplicate rows: {int(df.duplicated().sum())}")
doc.add_paragraph(f"Duplicate date-hour keys: {int(df.duplicated(['Date','Hour']).sum())}")

doc.add_heading("3. Key Findings", level=1)
doc.add_paragraph(f"Total rentals: {int(rentals.sum()):,}")
doc.add_paragraph(f"Average hourly rentals: {rentals.mean():.2f}")
doc.add_paragraph(f"Median hourly rentals: {rentals.median():.2f}")
doc.add_paragraph(
    f"Maximum observed rentals: {int(peak['Rented_Bike_Count']):,} on "
    f"{peak['Date'].date()} at {int(peak['Hour']):02d}:00."
)
doc.add_paragraph(
    f"Highest average demand hour: {int(hourly.idxmax()):02d}:00 "
    f"({hourly.max():.2f} rentals/hour)."
)
doc.add_heading("Average Rentals by Season", level=2)
table = doc.add_table(rows=1, cols=2)
table.style = "Light Shading Accent 1"
table.rows[0].cells[0].text = "Season"
table.rows[0].cells[1].text = "Average rentals/hour"
for season, value in seasonal.items():
    cells = table.add_row().cells
    cells[0].text = str(season)
    cells[1].text = f"{value:.2f}"

doc.add_heading("4. Charts", level=1)
for filename, caption in [
    ("hourly_demand.png", "Average demand by hour"),
    ("seasonal_demand.png", "Average demand by season"),
    ("rental_distribution.png", "Rental-count distribution"),
    ("temperature_vs_rentals.png", "Temperature versus rentals"),
    ("correlation_heatmap.png", "Correlation heatmap")
]:
    image_path = RESULTS / filename
    if image_path.exists():
        doc.add_paragraph(caption)
        doc.add_picture(str(image_path), width=Inches(5.8))

doc.add_heading("5. Interpretation and Limitations", level=1)
doc.add_paragraph(
    "Hourly and seasonal demand differ substantially. Temperature is positively associated "
    "with rentals, but correlation does not prove causation. Potential outliers should be "
    "reviewed in context before deciding whether to remove any observations. Results from "
    "this dataset may not generalise to other cities or periods."
)
doc.add_heading("Reference", level=1)
doc.add_paragraph(
    "UCI Machine Learning Repository, Seoul Bike Sharing Demand. "
    "https://archive.ics.uci.edu/dataset/560/seoul%2Bbike%2Bsharing%2Bdemand"
)
doc.save(REPORTS / "Task3_Real_World_Dataset_Analysis_Report.docx")

description = """This project analyses the Seoul Bike Sharing Demand dataset from the UCI Machine Learning Repository for the YuvaIntern Data Analyst Internship, Task 3.

The goal was to inspect the dataset, assess quality, explore hourly and seasonal demand, analyse numerical relationships, and communicate insights using Python, Pandas, NumPy, Matplotlib, and Jupyter Notebook. The dataset contains 8,760 hourly observations and 14 source attributes, including rental count, date, hour, temperature, humidity, wind speed, visibility, dew-point temperature, solar radiation, rainfall, snowfall, season, holiday status, and system functioning status.

Data quality checks covered missing values, exact duplicate rows, duplicate date-hour keys, data types, descriptive statistics, and domain validity. The cleaned dataset retained all 8,760 records. No missing values, exact duplicate rows, or duplicate date-hour keys were found.

The analysis identified 6,172,314 total rentals and average hourly demand of approximately 704.6 rentals. Average demand was highest at 18:00, reaching approximately 1,502.9 rentals per hour. Seasonal analysis showed summer had the highest average demand at approximately 1,034.1 rentals per hour, while winter had the lowest at approximately 225.5. Weekday average demand was approximately 719.5 rentals per hour, compared with 667.3 on weekends. Temperature had a positive correlation with rental demand; this relationship is an association and does not establish causation.

The submission includes a Jupyter Notebook, a Word report, cleaned data, quality checks, summary tables, and visualisations. Potential outliers are flagged for review rather than automatically removed. Dataset source: UCI Machine Learning Repository, Seoul Bike Sharing Demand. Source attribution and CC BY 4.0 licensing information are acknowledged.

GitHub repository: https://github.com/SommayDewat/YuvaIntern-Data-Analyst-Internship
"""
(REPORTS / "Task3_Portal_Submission_Description.txt").write_text(description, encoding="utf-8")

readme = """# Task 3 — Real-World Dataset Analysis

Dataset: Seoul Bike Sharing Demand, UCI Machine Learning Repository  
Source: https://archive.ics.uci.edu/dataset/560/seoul%2Bbike%2Bsharing%2Bdemand  
License: CC BY 4.0; retain attribution.

## Key findings
- 8,760 hourly records and 6,172,314 total rentals.
- Average hourly rentals: 704.6.
- Average demand peaks at 18:00.
- Summer has the highest seasonal average; winter has the lowest.
- Temperature is positively correlated with rentals; correlation does not prove causation.
- No missing values or duplicate rows were found.
- Potential outliers should be reviewed before removal.

## Structure
- `data/`: source dataset
- `notebooks/`: analysis notebook
- `src/`: scripts
- `results/`: cleaned dataset, summaries, and charts
- `reports/`: Word report and portal description
"""
(ROOT / "README.md").write_text(readme, encoding="utf-8")

print("DELIVERABLES CREATED")
print("Notebook:", NOTEBOOKS / "01_seoul_bike_demand_analysis.ipynb")
print("Word report:", REPORTS / "Task3_Real_World_Dataset_Analysis_Report.docx")
print("Portal description:", REPORTS / "Task3_Portal_Submission_Description.txt")
print("Task README:", ROOT / "README.md")
print("Portal description word count:", len(description.split()))