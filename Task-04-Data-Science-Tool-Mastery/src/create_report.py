
from pathlib import Path
import json

import pandas as pd
from docx import Document
from docx.shared import Inches, Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH

PROJECT_DIR = Path(__file__).resolve().parent.parent
DATA_DIR = PROJECT_DIR / "data"
RESULTS_DIR = PROJECT_DIR / "results"
CHART_DIR = RESULTS_DIR / "charts"
REPORT_DIR = PROJECT_DIR / "reports"

REPORT_DIR.mkdir(parents=True, exist_ok=True)

df = pd.read_csv(DATA_DIR / "restaurant_tips_cleaned.csv")

with (RESULTS_DIR / "analysis_summary.json").open(
    encoding="utf-8"
) as file:
    summary = json.load(file)

doc = Document()

# Document styling
styles = doc.styles
styles["Normal"].font.name = "Arial"
styles["Normal"].font.size = Pt(10)

title = doc.add_heading("Data Science Tool Mastery Project", 0)
title.alignment = WD_ALIGN_PARAGRAPH.CENTER

subtitle = doc.add_paragraph(
    "Exploratory Data Analysis and Visualization using Python"
)
subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER

doc.add_paragraph("Prepared by: Sommay Dewat")
doc.add_paragraph("Project Type: Data Science Tools Practical Project")

doc.add_heading("1. Executive Summary", level=1)
doc.add_paragraph(
    "This project demonstrates practical use of Python data science "
    "libraries for dataset inspection, cleaning, statistical analysis, "
    "feature engineering, and visualization. The Seaborn restaurant tips "
    "dataset was used to study bill amounts, tips, party sizes, days, "
    "and meal times. The original dataset contained 244 records and "
    "7 columns. One exact duplicate record was removed, leaving 243 "
    "records for analysis."
)

doc.add_heading("2. Project Objectives", level=1)
objectives = [
    "Inspect dataset structure, data types, distributions, and quality.",
    "Use Pandas to clean and transform data.",
    "Apply NumPy for numerical and statistical calculations.",
    "Create visualizations using Matplotlib and Seaborn.",
    "Communicate analytical findings and limitations.",
    "Practice reproducible analysis using Jupyter Notebook."
]
for item in objectives:
    doc.add_paragraph(item, style="List Bullet")

doc.add_heading("3. Tools and Technologies", level=1)
tools = [
    "Python 3.11",
    "Pandas for data manipulation and grouping",
    "NumPy for numerical calculations",
    "Matplotlib and Seaborn for visualizations",
    "Jupyter Notebook for documented analysis",
    "python-docx for report generation"
]
for item in tools:
    doc.add_paragraph(item, style="List Bullet")

doc.add_heading("4. Dataset Description", level=1)
doc.add_paragraph(
    "The dataset is Seaborn's built-in tips dataset, intended here "
    "for educational analysis. Its seven fields are total_bill, tip, "
    "sex, smoker, day, time, and size. The dataset describes restaurant "
    "bill and tip observations along with customer and meal attributes. "
    "Its small size makes it suitable for learning common data analysis "
    "workflows, but it should not be treated as representative of all "
    "restaurants or customers."
)

doc.add_heading("5. Data Inspection and Cleaning", level=1)
doc.add_paragraph(
    f"The original dataset had 244 rows and 7 columns. Inspection found "
    f"{sum(summary['missing_values_in_original_data'].values())} missing "
    f"values across the original columns and "
    f"{summary['duplicates_removed']} exact duplicate row. The duplicate "
    f"was removed. No rows were removed for missing numeric values or "
    f"invalid numeric records. The final cleaned dataset contains "
    f"{summary['rows_after_cleaning']} rows."
)
doc.add_paragraph(
    "A derived feature, tip_percentage, was calculated as the tip divided "
    "by the total bill multiplied by 100. Unusual values were not "
    "automatically removed solely because they appeared extreme."
)

doc.add_heading("6. Statistical Analysis", level=1)
stats = [
    f"Average total bill: {summary['total_bill_mean']:.2f}",
    f"Median total bill: {summary['total_bill_median']:.2f}",
    f"Standard deviation of total bill: {summary['total_bill_std']:.2f}",
    f"Average tip: {summary['tip_mean']:.2f}",
    f"Median tip: {summary['tip_median']:.2f}",
    f"Overall tip-to-bill ratio: "
    f"{summary['tip_to_bill_ratio_percent']:.2f}%"
]
for item in stats:
    doc.add_paragraph(item, style="List Bullet")

day_summary = df.groupby("day", observed=False)["tip"].mean()
if not day_summary.empty:
    highest_day = day_summary.idxmax()
    lowest_day = day_summary.idxmin()
    doc.add_paragraph(
        f"Among the recorded days, {highest_day} had the highest average "
        f"tip ({day_summary.max():.2f}), while {lowest_day} had the "
        f"lowest average tip ({day_summary.min():.2f})."
    )

doc.add_heading("7. Visualizations", level=1)

chart_descriptions = {
    "bill_distribution.png":
        "Histogram showing the distribution of restaurant bill amounts.",
    "average_tip_by_day.png":
        "Comparison of average tip amounts across recorded days.",
    "bill_vs_tip.png":
        "Scatter plot exploring the relationship between bills and tips.",
    "tip_percentage_by_time.png":
        "Comparison of average tip percentage across meal times.",
    "correlation_heatmap.png":
        "Heatmap showing pairwise correlations among numerical variables."
}

for filename, description in chart_descriptions.items():
    chart_path = CHART_DIR / filename
    if chart_path.exists():
        doc.add_paragraph(description)
        doc.add_picture(str(chart_path), width=Inches(5.8))
        doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER

doc.add_heading("8. Key Findings and Business Interpretation", level=1)
findings = [
    f"The cleaned sample has an average bill of "
    f"{summary['total_bill_mean']:.2f} and an average tip of "
    f"{summary['tip_mean']:.2f}.",
    f"The aggregate tip-to-bill ratio is "
    f"{summary['tip_to_bill_ratio_percent']:.2f}%. This is a ratio "
    "of total tips to total bills, not the mean of individual percentages.",
    "The daily average-tip comparison can help demonstrate how grouped "
    "summaries support exploratory business analysis.",
    "The bill-versus-tip scatter plot helps inspect whether larger bills "
    "tend to accompany larger tips; a visible association does not prove "
    "that bill size causes a particular tip."
]
for item in findings:
    doc.add_paragraph(item, style="List Bullet")

doc.add_heading("9. Recommendations", level=1)
recommendations = [
    "Use larger and more representative datasets before making operational decisions.",
    "Explore tip percentage by party size, day, and meal time.",
    "Investigate unusual observations before deciding whether they are errors.",
    "Use clear charts and descriptive statistics to communicate findings.",
    "Keep the notebook, source code, cleaned data, and outputs together for reproducibility."
]
for item in recommendations:
    doc.add_paragraph(item, style="List Bullet")

doc.add_heading("10. Limitations", level=1)
doc.add_paragraph(
    "This is a small educational dataset, so results cannot automatically "
    "be generalized to all restaurants. The analysis is descriptive and "
    "does not establish causality. The dataset's currency and collection "
    "context are not established in this report, so monetary values are "
    "reported in dataset units. The analysis does not establish statistical "
    "significance or support causal claims."
)

doc.add_heading("11. Conclusion", level=1)
doc.add_paragraph(
    "The project demonstrates an end-to-end introductory data science "
    "workflow: inspecting data, checking quality, removing an exact "
    "duplicate, calculating numerical summaries, engineering a feature, "
    "creating visualizations, and documenting findings. These exercises "
    "provide a practical foundation for further work with larger datasets "
    "and more advanced statistical methods."
)

doc.add_heading("12. Project Outputs", level=1)
outputs = [
    "data/restaurant_tips.csv",
    "data/restaurant_tips_cleaned.csv",
    "results/dataset_inspection.txt",
    "results/analysis_summary.json",
    "results/charts/",
    "notebooks/01_data_science_tool_mastery.ipynb",
    "src/"
]
for item in outputs:
    doc.add_paragraph(item, style="List Bullet")

report_path = REPORT_DIR / "Task4_Data_Science_Tool_Mastery_Report.docx"
doc.save(report_path)

print("Word report created successfully!")
print(f"Report path: {report_path}")
